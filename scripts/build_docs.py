from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor as PRGB

ROOT = Path('/mnt/data/TruthLens_AAT')
DOCS = ROOT/'docs'
FIGS = ROOT/'figures'
DOCS.mkdir(exist_ok=True)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def build_report(path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)

    styles = doc.styles
    styles['Normal'].font.name = 'Arial'
    styles['Normal'].font.size = Pt(10.5)
    styles['Title'].font.name = 'Arial'
    styles['Title'].font.size = Pt(22)
    styles['Heading 1'].font.name = 'Arial'
    styles['Heading 1'].font.size = Pt(15)
    styles['Heading 1'].font.color.rgb = RGBColor(31, 78, 121)
    styles['Heading 2'].font.name = 'Arial'
    styles['Heading 2'].font.size = Pt(12)
    styles['Heading 2'].font.color.rgb = RGBColor(54, 96, 146)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_after = Pt(18)
    r = p.add_run('B.M.S. COLLEGE OF ENGINEERING, BENGALURU')
    r.bold = True; r.font.size = Pt(16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('DEPARTMENT OF MACHINE LEARNING'); r.bold=True; r.font.size=Pt(13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Academic Activity Task (AAT)').bold = True
    doc.add_paragraph('')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('TRUTHLENS AI'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor(31,78,121)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('A Comparative Study of Traditional and Transformer-Based NLP Models\nfor Explainable Fake News Detection')
    r.bold=True; r.font.size=Pt(15)
    doc.add_paragraph('')
    for label, value in [
        ('Course / Activity', 'NLP - Topic Exploration, Implementation, Validation and Presentation'),
        ('Student 1', '[Name / USN]'),
        ('Student 2', '[Name / USN]'),
        ('Class / Semester', '[Insert class and semester]'),
        ('Faculty', '[Insert faculty name]'),
        ('Academic Year', '2026-27'),
    ]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=p.add_run(f'{label}: '); rr.bold=True
        p.add_run(value)
    doc.add_page_break()

    add_heading(doc, '1. AAT Requirement Mapping', 1)
    add_body(doc, 'The activity sheet asks the team to select an NLP topic, study its real-world application and architecture, implement the technique, evaluate the model, compare approaches where applicable, and submit a report and presentation. TruthLens AI is designed around each of these requirements.')
    table = doc.add_table(rows=1, cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    headers=['Activity requirement','TruthLens implementation','Evidence in project']
    for i,h in enumerate(headers): set_cell_text(table.rows[0].cells[i],h,True); shade_cell(table.rows[0].cells[i],'D9EAF7')
    rows=[
        ('Choose an NLP topic','Fake news text classification','Problem statement and literature review'),
        ('Real-time application / API','News screening interface + REST API','Streamlit app and FastAPI endpoint'),
        ('Study architecture / workflow','TF-IDF baseline and DistilBERT pipeline','Architecture and workflow figures'),
        ('Implement selected NLP model','Fine-tuned DistilBERT','train_transformer.py'),
        ('Pre-processing and evaluation','Cleaning, deduplication, stratified split, metrics','preprocess.py and evaluate.py'),
        ('Compare approaches','TF-IDF + Logistic Regression vs DistilBERT','metrics.csv + comparison graph'),
        ('Interpret results','LIME local explanations','explain.py + HTML output'),
        ('Report and presentation','DOCX report and PPT','docs/'),
    ]
    for row in rows:
        cells=table.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i],v)

    add_heading(doc, '2. Abstract', 1)
    add_body(doc, 'TruthLens AI is an NLP-based research prototype for classifying news articles into two labels, FAKE and REAL, using a labelled news corpus. The study compares a traditional TF-IDF plus Logistic Regression approach with a fine-tuned DistilBERT transformer. The work covers data cleaning, duplicate removal, stratified train/validation/test splitting, model training, evaluation using accuracy, precision, recall and F1-score, confusion-matrix analysis, and local model interpretation using LIME. A Streamlit interface and FastAPI endpoint demonstrate how the classifier could be used as a first-pass screening tool in a newsroom, moderation workflow, or information triage system.')
    add_body(doc, 'The model is intentionally described as a classification system rather than as a universal fact-checker. Its prediction is learned from the label distribution and writing patterns in the training corpus, so independent verification from trusted sources would be required before any high-stakes decision.')

    add_heading(doc, '3. Problem Statement', 1)
    add_body(doc, 'Large volumes of online news and social content make manual screening difficult. The goal of this project is to build a reproducible NLP pipeline that can classify article text using both a conventional feature-based model and a transformer model, then provide a human-readable explanation for an individual prediction.')
    add_heading(doc, '3.1 Objectives', 2)
    add_bullets(doc, [
        'Understand a practical NLP text-classification problem and its real-world use.',
        'Clean and prepare a public labelled news dataset for supervised learning.',
        'Create a traditional TF-IDF + Logistic Regression baseline.',
        'Fine-tune a DistilBERT transformer for the same binary classification task.',
        'Evaluate both models on an untouched test set using multiple metrics.',
        'Use LIME to inspect local evidence behind individual predictions.',
        'Package the final pipeline in a simple web application and REST API.',
    ])

    add_heading(doc, '4. Real-World Application', 1)
    add_body(doc, 'A practical use case is first-pass screening of article text before publication, moderation, or manual review. A newsroom or content platform could submit an article to an internal service, receive a model score, and route uncertain examples to a human reviewer. The output should be treated as a screening signal rather than as a final factual decision.')
    add_heading(doc, '4.1 User Workflow', 2)
    add_body(doc, 'A user pastes an article into the interface. The application checks whether enough text is available, sends the text to the trained model, displays the predicted class and confidence, and exposes probability information. The same prediction logic can be accessed programmatically through the FastAPI endpoint.')

    add_heading(doc, '5. Literature Review', 1)
    lit = [
        ('Devlin et al. (2019)', 'BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding', 'Introduced the BERT pre-training approach and established transformer encoders as a strong foundation for language understanding tasks.'),
        ('Sanh et al. (2019)', 'DistilBERT: A Distilled Version of BERT', 'Presented a smaller distilled transformer designed to retain much of BERT’s language understanding ability with lower computational cost.'),
        ('Shu et al. (2017)', 'Fake News Detection on Social Media: A Data Mining Perspective', 'Discussed fake-news detection as a broader data-mining problem and highlighted the role of content and contextual signals.'),
        ('Ribeiro et al. (2016)', 'Why Should I Trust You? Explaining the Predictions of Any Classifier', 'Introduced LIME, a model-agnostic approach for producing local explanations around individual predictions.'),
    ]
    table=doc.add_table(rows=1, cols=3); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(['Reference','Work','Relevance to project']): set_cell_text(table.rows[0].cells[i],h,True); shade_cell(table.rows[0].cells[i],'D9EAF7')
    for row in lit:
        cells=table.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i],v)
    add_body(doc, 'The literature supports the design choice of using a conventional baseline for comparison, a compact transformer for contextual language representation, and a local explainability method for inspection of individual predictions.')

    add_heading(doc, '6. Dataset', 1)
    add_body(doc, 'The project uses the ISOT Fake News Dataset provided by the Information Security and Object Technology Research Group at the University of Victoria. The dataset is distributed as separate True.csv and Fake.csv files. The project keeps only the article text and the binary label for the primary experiment.')
    add_body(doc, 'The original dataset is not a generic measure of truth on the open web. Its labels are tied to how the source corpus was assembled, so model performance can be influenced by publisher, time-period, topic and writing-style differences. This is addressed as a major limitation in the discussion.')
    add_heading(doc, '6.1 Data Preparation', 2)
    add_bullets(doc, [
        'Remove missing article text.',
        'Remove URLs from article text.',
        'Remove obvious Reuters prefix artifacts where present.',
        'Normalize repeated whitespace.',
        'Remove duplicate articles.',
        'Sample the same number of articles from each class for the primary experiment.',
        'Perform a stratified 70/15/15 train, validation and test split using seed 42.',
    ])
    add_body(doc, 'The default experiment uses 5,000 articles from each class, for 10,000 articles total before splitting. This keeps the project substantial while reducing training time for the academic prototype.')

    add_heading(doc, '7. System Architecture', 1)
    doc.add_picture(str(FIGS/'architecture.png'), width=Inches(6.9))
    p=doc.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_body(doc, 'Figure 1. Two-branch NLP architecture. The same cleaned article text is sent through a traditional TF-IDF baseline and a transformer-based DistilBERT branch. Both branches produce a binary prediction for comparison.')
    doc.add_picture(str(FIGS/'workflow.png'), width=Inches(6.9))
    p=doc.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_body(doc, 'Figure 2. End-to-end workflow from dataset ingestion to evaluation, explainability and application.')

    add_heading(doc, '8. Baseline Model: TF-IDF + Logistic Regression', 1)
    add_body(doc, 'The baseline uses a TF-IDF vectorizer with unigram and bigram features, English stop-word removal, sublinear TF scaling and a capped vocabulary of 20,000 features. Logistic Regression then learns a binary decision boundary over the sparse TF-IDF representation. This baseline is intentionally simple, fast, and useful for comparison.')
    add_heading(doc, '8.1 Why keep the baseline?', 2)
    add_bullets(doc, ['It establishes a reference point for the transformer.', 'It provides an inexpensive model for debugging the data pipeline.', 'It makes the comparison more informative than reporting only one model.', 'Its feature representation is easier to inspect than a deep transformer.'])

    add_heading(doc, '9. Transformer Model: DistilBERT', 1)
    add_body(doc, 'The second model is DistilBERT, a compact transformer encoder. The project uses the pretrained distilbert-base-uncased checkpoint as the language representation and replaces the final task head with a two-class classification layer. The model is fine-tuned on the labelled training split.')
    add_heading(doc, '9.1 Data Flow', 2)
    add_bullets(doc, [
        'Raw article text enters the tokenizer.',
        'The tokenizer converts text into subword token IDs and attention masks.',
        'The transformer encoder produces contextual hidden representations.',
        'The sequence representation is passed to the classification head.',
        'Softmax produces the two class probabilities.',
        'The class with the larger probability is returned as the prediction.',
    ])
    add_heading(doc, '9.2 Training Setup', 2)
    setup=[('Base checkpoint','distilbert-base-uncased'),('Maximum sequence length','256 tokens'),('Batch size','8'),('Learning rate','2e-5'),('Epochs','2'),('Optimizer','AdamW'),('Gradient clipping','1.0'),('Random seed','42')]
    table=doc.add_table(rows=1, cols=2); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(['Setting','Value']): set_cell_text(table.rows[0].cells[i],h,True); shade_cell(table.rows[0].cells[i],'D9EAF7')
    for k,v in setup:
        cells=table.add_row().cells; set_cell_text(cells[0],k); set_cell_text(cells[1],v)

    add_heading(doc, '10. Experimental Design', 1)
    add_body(doc, 'The dataset is split once using a fixed seed and the same partitions are reused by both models. The training split is used for parameter fitting, the validation split is used during transformer development and checkpoint selection, and the test split is held back for final comparison.')
    add_heading(doc, '10.1 Evaluation Metrics', 2)
    add_bullets(doc, [
        'Accuracy = correct predictions divided by all predictions.',
        'Precision = proportion of predicted positive examples that are positive.',
        'Recall = proportion of actual positive examples that are recovered.',
        'F1-score = harmonic mean of precision and recall.',
    ])

    add_heading(doc, '11. Results', 1)
    add_body(doc, 'The table below is generated from the actual saved test predictions. The placeholders must be replaced by the pipeline after the full experiment has been run.')
    table=doc.add_table(rows=1, cols=6); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(['Model','Accuracy','Precision','Recall','F1','Test size']): set_cell_text(table.rows[0].cells[i],h,True); shade_cell(table.rows[0].cells[i],'D9EAF7')
    vals=[('TF-IDF + Logistic Regression','[[LR_ACC]]','[[LR_PREC]]','[[LR_REC]]','[[LR_F1]]','[[TEST_SIZE]]'),('DistilBERT','[[BERT_ACC]]','[[BERT_PREC]]','[[BERT_REC]]','[[BERT_F1]]','[[TEST_SIZE]]')]
    for row in vals:
        cells=table.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i],v)
    add_body(doc, 'The most important comparison is not simply the highest accuracy. Precision, recall and F1 should be interpreted together, and the limitations of the dataset must be considered before drawing broad conclusions about real-world misinformation.')

    add_heading(doc, '12. Explainability with LIME', 1)
    add_body(doc, 'LIME creates a local explanation around one prediction by perturbing the text and fitting a simple surrogate model to approximate the classifier in that local neighbourhood. In TruthLens AI, the explanation is displayed as an inspection aid for a single article.')
    add_body(doc, 'The explanation is not presented as causal evidence that a particular word makes an article false. It only shows which parts of the input were influential for the trained model in that local example.')
    add_heading(doc, '12.1 Example Interpretation', 2)
    add_body(doc, 'During the demonstration, choose one held-out article, run explain.py, open results/lime_example.html, and discuss several highlighted words. Compare the explanation with the model class and confidence. The exact highlighted terms should be taken from the actual generated explanation rather than invented in the report.')

    add_heading(doc, '13. Application Prototype', 1)
    add_body(doc, 'The Streamlit interface exposes a small front end for the trained model. It accepts article text, performs tokenization, returns the predicted class and confidence, and displays the class probabilities. A FastAPI service exposes the same prediction logic through a POST /predict endpoint for integration into another application.')
    add_heading(doc, '13.1 Example API Request', 2)
    add_body(doc, '{"text": "Paste a sufficiently long article here"}')
    add_heading(doc, '13.2 Example API Response', 2)
    add_body(doc, '{"label": "REAL", "confidence": 0.xx, "probabilities": {"FAKE": 0.xx, "REAL": 0.xx}, "model": "DistilBERT"}')

    add_heading(doc, '14. Limitations and Responsible Use', 1)
    add_bullets(doc, [
        'The label source reflects the construction of the ISOT corpus rather than a universal truth database.',
        'Publisher, writing style, time period and topic can create shortcuts that are not equivalent to factual reasoning.',
        'The maximum sequence length truncates long articles, so information outside the first 256 tokens can be lost.',
        'LIME explanations are local and approximate rather than a proof of causation.',
        'A production system should add evidence retrieval, source verification, calibration and an explicit human-review path.',
    ])

    add_heading(doc, '15. Conclusion', 1)
    add_body(doc, 'TruthLens AI provides a complete NLP activity pipeline from data preparation through modelling, comparison, interpretation and application. The study demonstrates why a baseline remains useful when introducing a transformer model, and why evaluation and limitations are as important as the classifier itself. The project is intentionally reproducible: all data processing, training, evaluation and application steps are stored in one folder and can be repeated in Google Colab or a local Python environment.')

    add_heading(doc, '16. References', 1)
    refs=[
        'Devlin, J., Chang, M.-W., Lee, K., and Toutanova, K. (2019). BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding. NAACL-HLT.',
        'Sanh, V., Debut, L., Chaumond, J., and Wolf, T. (2019). DistilBERT, a distilled version of BERT: smaller, faster, cheaper and lighter. arXiv:1910.01108.',
        'Shu, K., Sliva, A., Wang, S., Tang, J., and Liu, H. (2017). Fake News Detection on Social Media: A Data Mining Perspective. SIGKDD Explorations.',
        'Ribeiro, M. T., Singh, S., and Guestrin, C. (2016). Why Should I Trust You? Explaining the Predictions of Any Classifier. KDD.',
        'ISOT Research Lab, University of Victoria. Fake News Detection Dataset. https://onlineacademiccommunity.uvic.ca/isot/2022/11/27/fake-news-detection-datasets/',
    ]
    add_bullets(doc, refs)

    add_heading(doc, 'Appendix A. Project File Map', 1)
    add_body(doc, 'app.py - Streamlit front end\napi.py - FastAPI service\nsrc/preprocess.py - data cleaning and split creation\nsrc/train_baseline.py - TF-IDF + Logistic Regression\nsrc/train_transformer.py - DistilBERT fine-tuning\nsrc/evaluate.py - held-out test evaluation\nsrc/explain.py - LIME explanation\nscripts/download_data.py - official dataset downloader\nscripts/smoke_test.py - local baseline smoke test\nscripts/build_submission_docs.py - metric placeholder replacement')

    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])

    doc.save(path)


def build_ppt(path):
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)

    def add_title(slide, title, subtitle=None):
        tx=slide.shapes.add_textbox(PInches(0.65), PInches(0.42), PInches(12), PInches(0.7))
        p=tx.text_frame.paragraphs[0]; p.text=title; p.font.size=PPt(27); p.font.bold=True; p.font.color.rgb=PRGB(31,78,121)
        if subtitle:
            tb=slide.shapes.add_textbox(PInches(0.68), PInches(1.05), PInches(11.8), PInches(0.45))
            pp=tb.text_frame.paragraphs[0]; pp.text=subtitle; pp.font.size=PPt(13); pp.font.color.rgb=PRGB(90,90,90)

    def add_bullets(slide, items, x=0.8, y=1.55, w=11.7, h=5.3, size=20):
        box=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h))
        tf=box.text_frame; tf.word_wrap=True; tf.clear()
        for i,item in enumerate(items):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=item; p.font.size=PPt(size); p.space_after=PPt(10); p.level=0
        return box

    def new_slide(title, subtitle=None):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb=PRGB(250,251,253)
        add_title(slide,title,subtitle)
        return slide

    # 1
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=PRGB(245,248,252)
    tb=s.shapes.add_textbox(PInches(0.85),PInches(1.25),PInches(11.6),PInches(2.2)); tf=tb.text_frame
    p=tf.paragraphs[0]; p.text='TruthLens AI'; p.font.size=PPt(40); p.font.bold=True; p.font.color.rgb=PRGB(31,78,121); p.alignment=PP_ALIGN.CENTER
    p=tf.add_paragraph(); p.text='Explainable Fake News Detection using NLP'; p.font.size=PPt(25); p.font.bold=True; p.alignment=PP_ALIGN.CENTER
    p=tf.add_paragraph(); p.text='AAT: Topic Exploration, Implementation, Validation and Presentation'; p.font.size=PPt(17); p.alignment=PP_ALIGN.CENTER
    p=s.shapes.add_textbox(PInches(2.0),PInches(4.6),PInches(9.3),PInches(1.1)); p.text_frame.text='[Student 1 / USN]    |    [Student 2 / USN]\nDepartment of Machine Learning, B.M.S. College of Engineering'
    for pp in p.text_frame.paragraphs: pp.font.size=PPt(16); pp.alignment=PP_ALIGN.CENTER

    # 2
    s=new_slide('1. Problem Statement','Why this is an NLP problem')
    add_bullets(s,['Online news creates a large volume of text that can be difficult to review manually.','The project studies whether article wording and contextual patterns can be used for supervised text classification.','The target is a screening signal that can support human review, not an automatic truth verdict.'])

    # 3
    s=new_slide('2. Objectives','What we implement')
    add_bullets(s,['Clean and validate a labelled news corpus.','Build a TF-IDF + Logistic Regression reference model.','Fine-tune DistilBERT for the same binary task.','Evaluate accuracy, precision, recall and F1 on the same held-out test set.','Use LIME for local explanations and expose the result through a simple application.'])

    # 4
    s=new_slide('3. Literature Review','Research basis')
    add_bullets(s,['BERT (Devlin et al., 2019): contextual transformer language representations.','DistilBERT (Sanh et al., 2019): compact transformer with lower computational cost.','Fake-news detection literature: content and contextual signals can be used in classification systems.','LIME (Ribeiro et al., 2016): local, model-agnostic explanations for individual predictions.'],size=18)

    # 5
    s=new_slide('4. Dataset and Pre-processing','ISOT Fake News Dataset')
    add_bullets(s,['Two source files: True.csv and Fake.csv.','Primary experiment: 5,000 examples per class, sampled with seed 42.','Cleaning: missing values, URLs, obvious source-prefix artifacts, whitespace and duplicates.','Split: 70% train, 15% validation and 15% test, stratified by label.'])

    # 6
    s=new_slide('5. System Architecture','Two modelling branches from the same cleaned text')
    s.shapes.add_picture(str(FIGS/'architecture.png'), PInches(1.0), PInches(1.6), width=PInches(11.3))

    # 7
    s=new_slide('6. Baseline: TF-IDF + Logistic Regression','Fast and interpretable reference')
    add_bullets(s,['TF-IDF converts words and bigrams into weighted sparse features.','Vocabulary capped at 20,000 features with sublinear term-frequency scaling.','Logistic Regression learns a binary decision boundary over the TF-IDF vectors.','This model helps separate genuine transformer gains from gains caused simply by preprocessing.'],size=18)

    # 8
    s=new_slide('7. Transformer: DistilBERT','Contextual language representation')
    add_bullets(s,['Tokenizer converts text to subword IDs and attention masks.','DistilBERT encoder builds contextual representations.','A two-class classification head produces FAKE and REAL logits.','Training uses AdamW, learning rate 2e-5, maximum length 256, batch size 8 and 2 epochs.'],size=18)

    # 9
    s=new_slide('8. Experimental Setup','Fair comparison')
    add_bullets(s,['Same cleaned text and same stratified test set for both models.','Fixed random seed: 42.','Primary evaluation: accuracy, precision, recall and F1.','Validation is used during transformer development; test data is held back until final evaluation.'],size=18)

    # 10
    s=new_slide('9. Results','Measured on the held-out test split after the final run')
    rows=[('Model','Accuracy','Precision','Recall','F1'),('TF-IDF + LR','[[LR_ACC]]','[[LR_PREC]]','[[LR_REC]]','[[LR_F1]]'),('DistilBERT','[[BERT_ACC]]','[[BERT_PREC]]','[[BERT_REC]]','[[BERT_F1]]')]
    table=s.shapes.add_table(3,5,PInches(0.55),PInches(1.8),PInches(6.55),PInches(2.25)).table
    widths=[1.75,1.2,1.2,1.2,1.2]
    for i,w in enumerate(widths): table.columns[i].width=PInches(w)
    for r,row in enumerate(rows):
        for c,val in enumerate(row):
            cell=table.cell(r,c); cell.text=val
            p=cell.text_frame.paragraphs[0]; p.font.size=PPt(13); p.alignment=PP_ALIGN.CENTER; p.font.bold=(r==0)
    add_bullets(s,['Interpret all four metrics together. Results are specific to this dataset and held-out test split.'],x=7.35,y=1.8,w=5.15,h=2.0,size=15)

    # 11
    s=new_slide('10. Explainability + Application','From model output to a usable prototype')
    add_bullets(s,['LIME creates a local explanation for an individual prediction.','Streamlit provides a simple article-analysis UI.','FastAPI exposes POST /predict for programmatic use.','The UI warns that the output is a learned classification signal, not a fact verdict.'],x=0.75,y=1.55,w=6.2,h=4.8,size=17)

    # 12
    s=new_slide('11. Limitations and Conclusion','What the project demonstrates')
    add_bullets(s,['Dataset source and publisher effects may create shortcuts that do not generalize to new sources.','Long articles are truncated to the first 256 tokens in the transformer experiment.','LIME is local and approximate.','The completed pipeline satisfies the AAT requirements: topic study, application, architecture, implementation, evaluation, comparison, interpretation, report and presentation.'],size=18)

    prs.save(path)


if __name__=='__main__':
    FIGS.mkdir(exist_ok=True)
    build_report(DOCS/'TruthLens_AAT_Report_Template.docx')
    build_ppt(DOCS/'TruthLens_AAT_Presentation_Template.pptx')
    print('Created report and presentation templates.')
