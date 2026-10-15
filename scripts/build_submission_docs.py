"""Build submission DOCX/PPTX from the measured experiment results."""

from pathlib import Path
import json

from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "metrics.json"
PROCESSED = ROOT / "data" / "processed" / "data_info.json"
DOCS = ROOT / "docs"

DOCX_TEMPLATE = DOCS / "TruthLens_AAT_Report_Template.docx"
PPTX_TEMPLATE = DOCS / "TruthLens_AAT_Presentation_Template.pptx"


def fmt(value):
    return f"{value * 100:.2f}%" if value is not None else "Not available"


def load_values():
    if not RESULTS.exists():
        raise FileNotFoundError("Run the full experiment first so results/metrics.json exists.")

    metrics = json.loads(RESULTS.read_text(encoding="utf-8"))
    data_info = {}
    if PROCESSED.exists():
        data_info = json.loads(PROCESSED.read_text(encoding="utf-8"))

    if data_info.get("sample_run") or metrics.get("sample_run"):
        raise ValueError(
            "The current results came from the smoke-test dataset. "
            "Run the real ISOT dataset before generating submission documents."
        )

    lr = metrics.get("logistic_regression_tfidf") or {}
    bert = metrics.get("distilbert") or {}
    return {
        "[[LR_ACC]]": fmt(lr.get("accuracy")),
        "[[LR_PREC]]": fmt(lr.get("precision")),
        "[[LR_REC]]": fmt(lr.get("recall")),
        "[[LR_F1]]": fmt(lr.get("f1")),
        "[[BERT_ACC]]": fmt(bert.get("accuracy")),
        "[[BERT_PREC]]": fmt(bert.get("precision")),
        "[[BERT_REC]]": fmt(bert.get("recall")),
        "[[BERT_F1]]": fmt(bert.get("f1")),
        "[[TEST_SIZE]]": str(metrics.get("dataset", {}).get("test_size", "N/A")),
    }


def replace_runs(paragraph, replacements):
    full = "".join(run.text for run in paragraph.runs)
    new = full
    for old, value in replacements.items():
        new = new.replace(old, value)
    if new != full and paragraph.runs:
        paragraph.runs[0].text = new
        for run in paragraph.runs[1:]:
            run.text = ""


def _insert_paragraph_after(paragraph):
    new_p = paragraph._p.__class__()
    paragraph._p.addnext(new_p)
    return paragraph._parent.add_paragraph()


def add_image_section(doc, title, image_path, width=5.8):
    if not image_path.exists():
        return False
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    p.alignment = 1
    p.add_run().add_picture(str(image_path), width=Inches(width))
    return True


def replace_docx(replacements):
    doc = Document(DOCX_TEMPLATE)
    for paragraph in doc.paragraphs:
        replace_runs(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_runs(paragraph, replacements)

    doc.add_heading("11.1 Evaluation Figures", level=2)
    add_image_section(doc, "11.1.1 Confusion Matrix", ROOT / "results" / "confusion_matrix.png")
    add_image_section(doc, "11.1.2 Model Comparison", ROOT / "results" / "model_comparison.png")
    add_image_section(doc, "11.1.3 Training Curves", ROOT / "results" / "training_curves.png")
    add_image_section(doc, "11.1.4 LIME Explanation", ROOT / "results" / "lime_example.png")

    out = DOCS / "TruthLens_AAT_Report_Submission.docx"
    doc.save(out)
    return out


def find_slide_by_title(prs, title_text):
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and title_text in shape.text:
                return slide
    return None


def replace_pptx(replacements):
    prs = Presentation(PPTX_TEMPLATE)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            for paragraph in shape.text_frame.paragraphs:
                replace_runs(paragraph, replacements)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            replace_runs(paragraph, replacements)

    results_slide = find_slide_by_title(prs, "9. Results")
    comparison = ROOT / "results" / "model_comparison.png"
    if results_slide:
        for shape in results_slide.shapes:
            if hasattr(shape, "text_frame"):
                if "Interpret all four metrics" in shape.text:
                    shape.text_frame.clear()
                elif "Replace placeholders with actual run output before submission" in shape.text:
                    shape.text_frame.paragraphs[0].text = "Measured on the held-out test split after the final run"

    if results_slide and comparison.exists():
        results_slide.shapes.add_picture(
            str(comparison), PInches(0.55), PInches(4.0), width=PInches(5.55)
        )

    explanation_slide = find_slide_by_title(prs, "10. Explainability + Application")
    lime = ROOT / "results" / "lime_example.png"
    if explanation_slide and lime.exists():
        explanation_slide.shapes.add_picture(
            str(lime), PInches(7.0), PInches(1.65), width=PInches(5.55)
        )

    out = DOCS / "TruthLens_AAT_Presentation_Submission.pptx"
    prs.save(out)
    return out


def main():
    replacements = load_values()
    report = replace_docx(replacements)
    ppt = replace_pptx(replacements)
    print(f"Created {report}")
    print(f"Created {ppt}")


if __name__ == "__main__":
    main()
